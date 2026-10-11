"""Builds Sethu_Presentation.pptx (16:9, editable text/shapes, embedded images)."""
import copy
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import art

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "assets"
GEN = ASSETS / "generated"
GEN.mkdir(parents=True, exist_ok=True)

FONT = "Segoe UI"


def rgb(c):
    return RGBColor(*c)


C_NAVY, C_NAVY2 = art.NAVY, art.NAVY2
C_BLUE, C_CYAN, C_RED, C_AMBER, C_WHITE, C_GREY = art.BLUE, art.CYAN, art.RED, art.AMBER, art.WHITE, art.GREY
C_GREEN = art.APP_GREEN
C_CARD = (13, 26, 56)
C_CARD_LINE = (38, 70, 130)

# ----------------------------------------------------------------------------------------------- assets
def make_assets():
    flood = art.flood_scene()
    flood.convert("RGB").save(GEN / "flood.jpg", quality=92)
    # title: flood with left-to-right navy gradient for text legibility
    t = flood.copy().convert("RGBA")
    ov = Image.new("RGBA", t.size, (0, 0, 0, 0))
    px = ov.load()
    for x in range(t.size[0]):
        a = int(235 * max(0, 1 - x / (t.size[0] * 0.78)) ** 1.2)
        for y in range(t.size[1]):
            px[x, y] = (5, 11, 26, a)
    t.alpha_composite(ov)
    t.convert("RGB").save(GEN / "title_bg.jpg", quality=92)
    art.darken(flood, 0.62).save(GEN / "flood_dark.jpg", quality=90)
    art.darken(flood, 0.30).save(GEN / "flood_mid.jpg", quality=90)
    art.dark_bg(accent=C_BLUE, glow_pos=(0.85, 0.1)).save(GEN / "bg_a.jpg", quality=92)
    art.dark_bg(accent=C_CYAN, glow_pos=(0.1, 0.9)).save(GEN / "bg_b.jpg", quality=92)
    art.dark_bg(accent=(120, 60, 200), glow_pos=(0.5, 0.0)).save(GEN / "bg_c.jpg", quality=92)
    for role, st in [("sender", "idle"), ("sender", "linked"), ("sender", "sent"), ("sender", "ended"),
                     ("relay", "listening"), ("relay", "linked"), ("relay", "held"), ("relay", "linked_coord"),
                     ("relay", "forwarding"), ("relay", "delivered"), ("coordinator", "listening"),
                     ("coordinator", "received"), ("coordinator", "connected")]:
        art.phone(role, st).save(GEN / f"phone_{role}_{st}.png")
    art.bt_badge(160).save(GEN / "bt_badge.png")
    art.bt_badge(160, ring=art.AMBER, fill=(60, 40, 10), active=False).save(GEN / "bt_badge_later.png")
    # three-device motif
    m = Image.new("RGBA", (900, 330), (0, 0, 0, 0))
    d = ImageDraw.Draw(m)
    xs = [110, 450, 790]
    for a, b in [(xs[0], xs[1]), (xs[1], xs[2])]:
        art.dashed_line(d, (a + 70, 165), (b - 70, 165), (53, 224, 255, 255), width=4, dash=14, gap=10)
    for x, role, st in zip(xs, ["sender", "relay", "coordinator"], ["sent", "held", "received"]):
        ph = art.phone(role, st, 0.26, label=False)
        m.paste(ph, (x - ph.size[0] // 2, 165 - ph.size[1] // 2), ph)
    m = m.convert("RGBA")
    m.alpha_composite(art.glow(m, 10, 0.7))
    d2 = ImageDraw.Draw(m)
    for x, role, st in zip(xs, ["sender", "relay", "coordinator"], ["sent", "held", "received"]):
        ph = art.phone(role, st, 0.26, label=False)
        m.paste(ph, (x - ph.size[0] // 2, 165 - ph.size[1] // 2), ph)
    for a, b in [(xs[0], xs[1]), (xs[1], xs[2])]:
        bb = art.bt_badge(60)
        m.paste(bb, ((a + b) // 2 - bb.size[0] // 2, 165 - bb.size[1] // 2), bb)
    m.save(GEN / "motif.png")
    # person icon
    p = Image.new("RGBA", (300, 300), (0, 0, 0, 0))
    pd = ImageDraw.Draw(p)
    pd.ellipse([105, 30, 195, 120], fill=(255, 255, 255, 255))
    pd.rounded_rectangle([70, 130, 230, 290], radius=60, fill=(255, 255, 255, 255))
    pd.line([(215, 150), (262, 70)], fill=(255, 255, 255, 255), width=26)
    p.save(GEN / "person.png")


# ----------------------------------------------------------------------------------------------- pptx helpers
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_alpha(shape, alpha_pct):
    sp = shape.fill._xPr.find(qn("a:solidFill"))
    clr = sp[0]
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha_pct * 1000)))


def new_slide(bg, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(str(GEN / bg), 0, 0, prs.slide_width, prs.slide_height)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def tb(slide, x, y, w, h, text, size=18, color=C_WHITE, bold=False, align="l", anchor="t", italic=False, spacing=None, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = ln
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = rgb(color)
    if name:
        box.name = name
    return box


def card(slide, x, y, w, h, fill=C_CARD, line=C_CARD_LINE, alpha=88, radius=0.12, lw=1.25, dash=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    set_alpha(sh, alpha)
    sh.line.color.rgb = rgb(line)
    sh.line.width = Pt(lw)
    if dash:
        sh.line.dash_style = 4  # dash
    sh.shadow.inherit = False
    return sh


def chip(slide, x, y, w, h, text, fg=C_CYAN, bg=(14, 40, 100), size=12, bold=True, line=None):
    sh = card(slide, x, y, w, h, fill=bg, line=line or fg, alpha=95, radius=0.5, lw=1)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(fg)
    return sh


def arrow(slide, x1, y1, x2, y2, color=C_CYAN, width=2.5, dash=False, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    if head:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle")
        te.set("w", "med")
        te.set("len", "med")
    return c


def pic(slide, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(GEN / name), Inches(x), Inches(y), **kw)


def header(slide, kicker, title, n):
    tb(slide, 0.6, 0.38, 9, 0.35, kicker.upper(), size=13, color=C_CYAN, bold=True)
    tb(slide, 0.6, 0.68, 12, 0.8, title, size=34, bold=True)
    tb(slide, 0.6, 7.06, 6, 0.3, "SETHU  |  Hackathon presentation", size=10, color=(110, 130, 165))
    tb(slide, 12.1, 7.06, 0.7, 0.3, f"{n}", size=10, color=(110, 130, 165), align="r")


def bullets_card(slide, x, y, w, h, title, lines, accent=C_CYAN, title_size=17, line_size=13.5):
    card(slide, x, y, w, h)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.18), Inches(0.07), Inches(h - 0.36))
    bar.fill.solid(); bar.fill.fore_color.rgb = rgb(accent); bar.line.fill.background()
    tb(slide, x + 0.22, y + 0.12, w - 0.35, 0.4, title, size=title_size, bold=True)
    tb(slide, x + 0.22, y + 0.55, w - 0.35, h - 0.6, lines, size=line_size, color=(190, 205, 230), spacing=3)


# ----------------------------------------------------------------------------------------------- slides
def build():
    # 1 ------------------------------------------------------------------------------- title
    s = new_slide("title_bg.jpg", "Open with the flood scenario. Sethu = Offline Emergency Communication & Disaster Response. "
                                  "Illustration only: the artwork is drawn, not a photograph.")
    chip(s, 0.7, 0.55, 3.35, 0.38, "HACKATHON PRESENTATION", size=12)
    tb(s, 0.62, 1.25, 8, 1.6, "SETHU", size=104, bold=True, name="Title")
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(2.95), Inches(2.6), Inches(0.07))
    ln.fill.solid(); ln.fill.fore_color.rgb = rgb(C_CYAN); ln.line.fill.background()
    tb(s, 0.7, 3.2, 7.4, 1.2, ["Offline Emergency Communication", "& Disaster Response"], size=30, bold=True)
    tb(s, 0.7, 4.55, 6.6, 1.0, "Connecting emergency reports when conventional connectivity is unavailable.", size=19, color=(200, 215, 240), italic=True)
    pic(s, "motif.png", 0.55, 5.25, w=5.4)
    card(s, 10.15, 0.55, 2.55, 0.95, fill=(60, 12, 20), line=C_RED, alpha=85, radius=0.2)
    tb(s, 10.15, 0.62, 2.55, 0.4, "NO NETWORK", size=17, bold=True, color=(255, 120, 130), align="c")
    tb(s, 10.15, 1.0, 2.55, 0.4, "conventional connectivity down", size=11, color=(230, 170, 175), align="c")
    tb(s, 7.3, 6.95, 5.6, 0.35, "Illustrative artwork - not a photograph", size=10, color=(150, 160, 190), align="r")

    # 2 ------------------------------------------------------------------------------- problem
    s = new_slide("flood_mid.jpg", "A believable scenario: a flood, three people on a terrace, one needs urgent medical help, one needs evacuation. "
                                   "A volunteer has a phone but the message cannot reach the response coordinator. Illustrative, not a claim about a real event.")
    header(s, "The problem", "When networks fail, reports stall", 2)
    tb(s, 0.6, 1.4, 6.2, 1.0, ["Three people on a terrace. Water on the lower floors.", "One needs a medic. One needs evacuation."],
       size=17, color=(225, 233, 250), spacing=2)
    items = [("Reports may not reach a coordinator", C_RED),
             ("Connectivity may be intermittent or unavailable", C_AMBER),
             ("Delayed reports make coordination harder", C_AMBER),
             ("Nearby phones can still talk to each other locally", C_CYAN)]
    for i, (t, c) in enumerate(items):
        y = 2.5 + i * 1.08
        card(s, 0.6, y, 5.9, 0.92, fill=(6, 12, 30), line=c, alpha=88, radius=0.18)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.82), Inches(y + 0.29), Inches(0.34), Inches(0.34))
        dot.fill.solid(); dot.fill.fore_color.rgb = rgb(c); dot.line.fill.background()
        tb(s, 1.3, y + 0.08, 5.1, 0.78, t, size=15.5, bold=True, anchor="m")
    tb(s, 7.3, 6.7, 5.5, 0.3, "Illustrative artwork - scenarios are examples, not real deployments", size=10, color=(170, 178, 205), align="r")

    # 3 ------------------------------------------------------------------------------- solution
    s = new_slide("bg_a.jpg", "Sethu moves an emergency report phone to phone using local Bluetooth. The Relay stores it, so a lost connection does not lose the report. "
                              "Mockup phones are illustrative, not screenshots.")
    header(s, "The solution", "Store on the relay. Forward when a link appears.", 3)
    nodes = [("Person in danger", None, "person.png"), ("Sender phone", "sender_sent", None),
             ("Relay phone", "relay_held", None), ("Coordinator phone", "coordinator_received", None)]
    xs = [0.9, 3.95, 7.0, 10.05]
    for (label, ph, img), x in zip(nodes, xs):
        card(s, x, 1.75, 2.55, 4.3, alpha=70)
        if img:
            p = pic(s, img, x + 0.4, 2.6, w=1.75)
        else:
            pic(s, f"phone_{ph}.png", x + 0.62, 1.95, h=3.3)
        tb(s, x, 5.38, 2.55, 0.5, label, size=17, bold=True, align="c")
    sub = ["", "creates the SOS", "stores it, then ACKs", "saves and shows it"]
    for t, x in zip(sub, xs):
        if t:
            tb(s, x, 5.7, 2.55, 0.35, t, size=12, color=C_GREY, align="c")
    tb(s, xs[0], 5.7, 2.55, 0.35, "needs help", size=12, color=C_GREY, align="c")
    for x in xs[:3]:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.6), Inches(3.55), Inches(0.4), Inches(0.5))
        a.fill.solid(); a.fill.fore_color.rgb = rgb(C_CYAN); a.line.fill.background()
    card(s, 0.9, 6.25, 11.7, 0.62, fill=(12, 36, 80), line=C_CYAN, alpha=90, radius=0.3)
    tb(s, 0.9, 6.3, 11.7, 0.5, "A stored report survives a temporary disconnection and goes onward when the next connection exists.",
       size=16, bold=True, align="c", anchor="m")

    # 4 ------------------------------------------------------------------------------- architecture (3 phones)
    s = new_slide("bg_b.jpg", "Roles are fixed by the user: iQOO 15 = Sender, Vivo Y15 = Relay, Motorola Edge 30 Pro = Coordinator. "
                              "The implementation keeps ONE active Bluetooth link per phone at a time, so Stage B only starts after Stage A is ended. "
                              "Connections are made by the user (Start Server, then tap the peer). Forwarding is automatic.")
    header(s, "Three-phone architecture", "Two stages, one active link at a time", 4)
    px = [1.55, 5.65, 9.75]
    labels = [("iQOO 15", "SOS Sender"), ("Vivo Y15", "Relay"), ("Motorola Edge 30 Pro", "Coordinator")]
    imgs = ["phone_sender_sent.png", "phone_relay_linked.png", "phone_coordinator_received.png"]
    for x, (n, r), im in zip(px, labels, imgs):
        pic(s, im, x, 2.05, h=3.5)
        tb(s, x - 0.55, 5.65, 3.0, 0.4, n, size=17, bold=True, align="c")
        tb(s, x - 0.55, 6.0, 3.0, 0.35, r, size=13, color=C_CYAN, align="c")
    # Stage A (solid)
    chip(s, 1.55, 1.38, 4.2, 0.4, "STAGE A  -  iQOO  <->  Vivo  (manual link)", size=12)
    arrow(s, px[0] + 1.85, 3.6, px[1] - 0.1, 3.6, color=C_CYAN, width=3, head=False)
    pic(s, "bt_badge.png", (px[0] + px[1]) / 2 + 0.55 - 0.45, 3.15, w=0.9)
    # Stage B (dashed: later)
    chip(s, 7.2, 1.38, 4.9, 0.4, "STAGE B  -  Vivo  <->  Motorola  (manual, after A ends)", fg=C_AMBER, bg=(52, 38, 10), size=12)
    arrow(s, px[1] + 1.85, 3.6, px[2] - 0.1, 3.6, color=C_AMBER, width=3, dash=True, head=False)
    pic(s, "bt_badge_later.png", (px[1] + px[2]) / 2 + 0.55 - 0.45, 3.15, w=0.9)
    card(s, 1.2, 6.55, 10.9, 0.5, fill=(10, 24, 55), line=C_CYAN, alpha=70, radius=0.5)
    tb(s, 1.2, 6.58, 10.9, 0.45, "You connect each pair manually.  Forwarding of the stored SOS is automatic once Stage B is up.",
       size=14, bold=True, color=C_WHITE, align="c")
    tb(s, 9.3, 0.4, 3.5, 0.3, "Illustrative mockups, not screenshots", size=10, color=(150, 160, 190), align="r")

    # 5 ------------------------------------------------------------------------------- workflow
    s = new_slide("bg_c.jpg", "Seven steps. Steps 1-3 happen on the first manual link. Step 4 ends it (for example iQOO Bluetooth off). "
                              "Step 5 is the user's second manual connection. Step 6 is automatic. Step 7: the Coordinator saves first, then acknowledges, and only then does the Relay mark delivered.")
    header(s, "Live SOS workflow", "From first tap to a confirmed delivery", 5)
    steps = [("1", "Create SOS", C_CYAN), ("2", "Send to Relay", C_CYAN), ("3", "Relay saves it, then ACKs", C_CYAN),
             ("4", "Original link ends", C_AMBER), ("5", "Connect to Coordinator (manual)", C_AMBER),
             ("6", "Forward automatically", C_GREEN), ("7", "Coordinator saves + ACKs", C_GREEN)]
    sw = 1.66
    for i, (n, t, c) in enumerate(steps):
        x = 0.6 + i * (sw + 0.08)
        card(s, x, 1.55, sw, 1.35, alpha=85, line=c, radius=0.12)
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.1), Inches(1.65), Inches(0.42), Inches(0.42))
        o.fill.solid(); o.fill.fore_color.rgb = rgb(c); o.line.fill.background()
        o.text_frame.margin_left = o.text_frame.margin_right = 0
        o.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = o.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
        r = pp.add_run(); r.text = n; r.font.size = Pt(15); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = rgb(C_NAVY)
        tb(s, x + 0.08, 2.12, sw - 0.14, 0.78, t, size=12.5, bold=True, color=C_WHITE)
    frames = [("phone_sender_sent.png", "phone_relay_held.png", "Steps 1-3", "Sender ACKed by Relay"),
              (None, "phone_relay_held.png", "Step 4", "Relay holds the pending SOS"),
              ("phone_relay_forwarding.png", "phone_coordinator_connected.png", "Steps 5-6", "Auto-forward, awaiting ACK"),
              ("phone_relay_delivered.png", "phone_coordinator_received.png", "Step 7", "Delivered only after ACK")]
    gx = [0.6, 3.6, 6.6, 9.6]
    for (a, b, tag, cap), x in zip(frames, gx):
        card(s, x, 3.15, 2.85, 3.65, alpha=60)
        tb(s, x, 3.2, 2.85, 0.3, tag, size=13, bold=True, color=C_CYAN, align="c")
        if a:
            pic(s, a, x + 0.1, 3.55, h=2.55)
            pic(s, b, x + 1.5, 3.55, h=2.55)
        else:
            pic(s, b, x + 0.8, 3.55, h=2.55)
        tb(s, x, 6.2, 2.85, 0.5, cap, size=12, color=(200, 212, 235), align="c")
    tb(s, 8.0, 0.42, 4.8, 0.3, "Illustrative mockups, not screenshots", size=10, color=(150, 160, 190), align="r")

    # 6 ------------------------------------------------------------------------------- reliability
    s = new_slide("bg_a.jpg", "These mechanisms exist in the code and are exercised by unit and simulated tests. "
                              "DELIVERED is only recorded after the Coordinator's acknowledgement for the same report ID on the same connection.")
    header(s, "Reliability and data integrity", "A report is never 'delivered' on a hunch", 6)
    mech = [("Durable local storage", "Written to disk before any ACK is sent"),
            ("Unique report IDs", "Same ID = same incident everywhere"),
            ("Duplicate prevention", "Retransmissions do not create copies"),
            ("Acknowledgement states", "Matching ID, session and stage required"),
            ("Pending queues", "Report waits on the Relay for a link"),
            ("Controlled retries", "Max 3 sends per connection; no auto-dialing")]
    for i, (t, d) in enumerate(mech):
        x = 0.6 + (i % 3) * 4.1
        y = 1.55 + (i // 3) * 1.45
        bullets_card(s, x, y, 3.9, 1.3, t, [d], accent=C_CYAN, title_size=16, line_size=13)
    tb(s, 0.6, 4.55, 6, 0.35, "RELAY DELIVERY STATES", size=12, bold=True, color=C_CYAN)
    chain = [("PENDING FORWARD", C_AMBER, (52, 38, 10)), ("FORWARDED  (awaiting ACK)", (110, 175, 255), (14, 34, 64)), ("DELIVERED TO COORDINATOR", C_GREEN, (14, 48, 34))]
    cx = 0.6
    widths = [2.3, 3.0, 3.3]
    for (t, fg, bg), w_ in zip(chain, widths):
        chip(s, cx, 4.92, w_, 0.55, t, fg=fg, bg=bg, size=13)
        if cx + w_ < 9.0:
            arrow(s, cx + w_ + 0.02, 5.2, cx + w_ + 0.34, 5.2, color=C_GREY, width=2.2)
        cx += w_ + 0.38
    chip(s, 10.25, 4.92, 2.5, 0.55, "FAILED: stays pending", fg=(255, 130, 140), bg=(56, 14, 22), size=11.5)
    tb(s, 0.6, 5.65, 6, 0.35, "SENDER DELIVERY STATES", size=12, bold=True, color=C_CYAN)
    chain2 = [("PENDING", C_AMBER, (52, 38, 10)), ("TRANSMITTED", (110, 175, 255), (14, 34, 64)),
              ("ACKED BY RELAY", C_GREEN, (14, 48, 34)), ("ACKED BY COORDINATOR", C_GREEN, (14, 48, 34))]
    cx = 0.6
    for (t, fg, bg), w_ in zip(chain2, [1.9, 2.4, 2.7, 3.2]):
        chip(s, cx, 6.02, w_, 0.5, t, fg=fg, bg=bg, size=12)
        if cx + w_ < 9.6:
            arrow(s, cx + w_ + 0.02, 6.27, cx + w_ + 0.3, 6.27, color=C_GREY, width=2.2)
        cx += w_ + 0.34
    tb(s, 0.6, 6.62, 12, 0.35, "Socket connected or write returned is never treated as proof of delivery.", size=13, color=(190, 205, 230), italic=True)

    # 7 ------------------------------------------------------------------------------- offline-first
    s = new_slide("bg_b.jpg", "Checked against the merged debug and release manifests: no INTERNET permission. "
                              "The local LLM weights and the offline map file are loaded on the device (not bundled in the APK). "
                              "Bluetooth pairing is not app-level authentication, and there is no app-level encryption yet.")
    header(s, "Offline-first engineering", "Everything the demo needs runs on the phone", 7)
    tiles = [("Bluetooth Classic", "RFCOMM link between nearby phones"),
             ("On-device storage", "Atomic JSON files for every report"),
             ("Offline voice", "Vosk speech model bundled in the app"),
             ("Local AI assist", "Qwen2.5-0.5B via llama.cpp; model file side-loaded"),
             ("Offline maps", "Mapsforge vector map file loaded on the device")]
    for i, (t, d) in enumerate(tiles):
        x = 0.6 + i * 2.45
        card(s, x, 1.65, 2.3, 2.25, alpha=85)
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(1.85), Inches(0.12), Inches(0.12))
        o.fill.solid(); o.fill.fore_color.rgb = rgb(C_CYAN); o.line.fill.background()
        tb(s, x + 0.15, 2.05, 2.05, 0.8, t, size=16, bold=True)
        tb(s, x + 0.15, 2.85, 2.05, 1.25, d, size=12.5, color=(190, 205, 230))
    card(s, 0.6, 4.2, 6.0, 1.45, fill=(8, 40, 36), line=C_GREEN, alpha=90)
    tb(s, 0.85, 4.3, 5.6, 0.4, "No INTERNET permission", size=19, bold=True, color=C_GREEN)
    tb(s, 0.85, 4.8, 5.6, 0.8, "Verified in the merged debug and release Android manifests. No cloud, analytics or telemetry dependency.",
       size=13, color=(190, 230, 215))
    card(s, 6.85, 4.2, 5.9, 1.45, fill=(52, 38, 10), line=C_AMBER, alpha=90)
    tb(s, 7.1, 4.3, 5.5, 0.4, "What this does not claim", size=19, bold=True, color=C_AMBER)
    tb(s, 7.1, 4.8, 5.5, 0.8, "Bluetooth pairing is not app-level authentication, and Sethu adds no encryption of its own yet.",
       size=13, color=(245, 225, 180))
    tb(s, 0.6, 5.95, 12, 0.6, "Offline-first for the intended demo workflow; this is not a security guarantee.", size=14, italic=True, color=(190, 205, 230))

    # 8 ------------------------------------------------------------------------------- coordination
    s = new_slide("bg_c.jpg", "The priority engine is deterministic: explicit keyword rules with a stated reason, not a trained model. "
                              "The local LLM only proposes structured fields that the person reviews; unknown stays unknown. "
                              "No diagnosis, dispatch or tracking is claimed.")
    header(s, "Intelligent emergency coordination", "Explainable priority, not a black box", 8)
    tb(s, 0.6, 1.5, 6.2, 0.35, "RULE-BASED PRIORITY ENGINE", size=12, bold=True, color=C_CYAN)
    levels = [("CRITICAL", "Explicit life-threatening wording", C_RED, (60, 12, 20)),
              ("HIGH", "Trapped, serious injury, fire, urgent medical", C_AMBER, (60, 40, 8)),
              ("MEDIUM", "Needs help, no life-threat evidence", (110, 175, 255), (14, 34, 64)),
              ("UNKNOWN", "Too little information to classify", C_GREY, (28, 36, 56))]
    for i, (n, d, fg, bg) in enumerate(levels):
        y = 1.95 + i * 0.95
        chip(s, 0.6, y, 1.7, 0.6, n, fg=fg, bg=bg, size=14)
        tb(s, 2.5, y + 0.05, 4.4, 0.55, d, size=14, color=(205, 218, 240), anchor="m")
    tb(s, 0.6, 5.85, 6.3, 0.9, ["Highest explicitly supported level wins.", "Each result carries a readable reason."], size=13, color=C_GREY, spacing=2)
    card(s, 7.2, 1.5, 5.5, 2.35, alpha=85)
    tb(s, 7.45, 1.6, 5.1, 0.4, "Local AI assist (optional)", size=18, bold=True)
    tb(s, 7.45, 2.1, 5.1, 1.7, ["Qwen2.5-0.5B on-device (llama.cpp)", "Proposes category and people count",
                                "Unknown stays unknown - never invented", "Proposal is editable; the person confirms it"],
       size=13.5, color=(200, 214, 238), spacing=3)
    card(s, 7.2, 4.05, 5.5, 2.7, alpha=85)
    pic(s, "phone_coordinator_received.png", 7.45, 4.15, h=2.5)
    tb(s, 9.2, 4.3, 3.4, 1.7, ["Coordinator dashboard:", "newest received report is highlighted;", "a priority filter never hides a new SOS."],
       size=13, color=(200, 214, 238), spacing=3)
    tb(s, 9.2, 6.2, 3.4, 0.5, "Illustrative mockup", size=10, color=(150, 160, 190))

    # 9 ------------------------------------------------------------------------------- tech architecture
    s = new_slide("bg_a.jpg", "Every box is a real class or library in the repository. Kotlin + Jetpack Compose. The transport is Bluetooth Classic RFCOMM through a thin adapter "
                              "that the simulated tests replace with fake sockets.")
    header(s, "Technical architecture", "What is actually in the repository", 9)

    def layer(y, h, title, boxes, accent, x0=0.6, w_total=12.1):
        card(s, x0, y, w_total, h, alpha=55, line=accent, radius=0.08)
        tb(s, x0 + 0.15, y + 0.04, 3, 0.3, title.upper(), size=10.5, bold=True, color=accent)
        n = len(boxes)
        gap = 0.14
        bw = (w_total - 0.3 - gap * (n - 1)) / n
        for i, t in enumerate(boxes):
            bx = x0 + 0.15 + i * (bw + gap)
            b = card(s, bx, y + 0.36, bw, h - 0.46, fill=(18, 36, 78), line=accent, alpha=95, radius=0.12, lw=1)
            tb(s, bx, y + 0.36, bw, h - 0.46, t, size=11.5, bold=True, align="c", anchor="m")

    for yy in (2.38, 3.47, 4.56, 5.65):
        arrow(s, 6.65, yy, 6.65, yy + 0.2, color=C_GREY, width=2.5)
    layer(1.5, 0.88, "Android UI - Jetpack Compose", ["SosSenderScreen", "RelayDeviceScreen", "CoordinatorScreen", "Direct Device Link (OfflineCommScreen)", "OfflineMapScreen"], C_CYAN)
    layer(2.6, 0.88, "On-device intelligence", ["PriorityClassifier (rules)", "OfflineSpeechManager (Vosk)", "OfflineAiExtractor (llama.cpp, Qwen2.5)", "MapsforgeMapView (offline maps)"], (150, 120, 255))
    layer(3.69, 0.88, "Delivery engine", ["StoreAndForwardController (link-driven, never dials)", "Role + ACK validation (Sender / Relay / Coordinator)"], C_GREEN)
    layer(4.78, 0.88, "Protocol and durable stores", ["SethuMessage + SosReport (JSON lines)", "SenderReportStore", "RelayReportStore", "CoordinatorReportStore"], C_AMBER)
    layer(5.87, 0.88, "Transport", ["BluetoothRfcommTransport", "RfcommAdapter  ->  AndroidRfcommAdapter", "Android Bluetooth Classic (RFCOMM)"], C_BLUE)
    tb(s, 8.6, 0.42, 4.2, 0.3, "Kotlin | Jetpack Compose | coroutines", size=11, color=(150, 170, 205), align="r")

    # 10 ------------------------------------------------------------------------------ results
    s = new_slide("bg_b.jpg", "Be explicit with the jury: unit tests and simulated end-to-end tests are verified. Physical Bluetooth validation is in progress. "
                              "The screenshot on the right is a REAL screenshot of the iQOO 15 from the first physical connection attempt, which failed and is being debugged. "
                              "Delete that box if you prefer not to show it, but do not claim physical success.")
    header(s, "Demonstration results", "What is verified - and what is not yet", 10)
    cols = [("Automated unit tests", "VERIFIED", C_GREEN, (14, 48, 34),
             ["188 unit tests, 0 failed.", "Covers stores, protocol, priority rules, ACK correlation and persistence."]),
            ("Simulated end-to-end", "VERIFIED (logic)", C_GREEN, (14, 48, 34),
             ["Five-stage workflow on fake Bluetooth sockets.", "Also lost ACKs, restarts, duplicates and connection flapping."]),
            ("Physical Bluetooth", "IN PROGRESS", C_AMBER, (52, 38, 10),
             ["Three-phone run not yet verified.", "First iQOO > Vivo attempt failed at the RFCOMM connect step; debugging."])]
    for i, (t, st, fg, bg, lines) in enumerate(cols):
        x = 0.6 + i * 3.0
        card(s, x, 1.6, 2.85, 3.95, alpha=85, line=fg)
        tb(s, x + 0.15, 1.72, 2.6, 0.7, t, size=16, bold=True)
        chip(s, x + 0.15, 2.5, 2.55, 0.42, st, fg=fg, bg=bg, size=12)
        tb(s, x + 0.15, 3.1, 2.6, 2.3, lines, size=13.5, color=(200, 214, 238), spacing=8)
    # real screenshot
    shot = Image.open(ASSETS / "real" / "iqoo_direct_link_physical_attempt.webp").convert("RGB")
    shot = shot.crop((0, 120, shot.size[0], shot.size[1]))
    shot.save(GEN / "real_iqoo_attempt.png")
    card(s, 9.75, 1.6, 3.0, 5.1, alpha=85, line=C_AMBER)
    pic(s, "real_iqoo_attempt.png", 10.25, 1.75, h=4.0)
    tb(s, 9.8, 5.85, 2.9, 0.8, ["REAL screenshot - iQOO 15", "first physical attempt (failed)"], size=11, color=C_AMBER, bold=True, align="c")
    card(s, 0.6, 5.75, 8.9, 0.95, fill=(12, 36, 80), line=C_CYAN, alpha=90, radius=0.2)
    tb(s, 0.8, 5.8, 8.5, 0.85, ["Intended physical demonstration: iQOO > Vivo, send SOS, iQOO link ends,",
                                "Vivo > Motorola (manual), automatic forward, delivered after ACK."], size=13.5, bold=True, anchor="m")
    tb(s, 0.6, 6.78, 9, 0.3, "Simulated tests validate application logic only, not the Bluetooth stack on the phones.", size=11, italic=True, color=C_GREY)

    # 11 ------------------------------------------------------------------------------ future
    s = new_slide("bg_c.jpg", "All of these are future improvements, not current capabilities.")
    header(s, "Impact and future scope", "Where this approach could go next", 11)
    fut = [("Multi-hop relays", "Reports across several relay phones"),
           ("Background delivery", "More robust delivery when the app is not open"),
           ("Authentication + encryption", "Message-level security"),
           ("Better prioritization", "Richer, evaluated triage rules"),
           ("More languages", "Accessible, multilingual interfaces"),
           ("Authorized integration", "Link to official emergency-response systems"),
           ("Field testing", "Larger tests in disaster-like conditions")]
    for i, (t, d) in enumerate(fut):
        col, row = i % 4, i // 4
        x = 0.6 + col * 3.05
        y = 1.65 + row * 2.45
        card(s, x, y, 2.9, 2.2, alpha=85)
        chip(s, x + 0.15, y + 0.15, 1.3, 0.34, "FUTURE", fg=C_AMBER, bg=(52, 38, 10), size=10.5)
        tb(s, x + 0.15, y + 0.65, 2.6, 0.7, t, size=16.5, bold=True)
        tb(s, x + 0.15, y + 1.3, 2.6, 0.85, d, size=12.5, color=(190, 205, 230))
    tb(s, 9.75, 4.2, 3.0, 2.2, ["Sethu supports", "emergency services.", "It does not replace them."], size=17, bold=True, color=C_CYAN, spacing=2)

    # 12 ------------------------------------------------------------------------------ closing
    s = new_slide("flood_dark.jpg", "Close on the tagline.")
    tb(s, 0.5, 1.2, 12.3, 1.6, ["When connectivity breaks, emergency communication", "should not stop at the first disconnected phone."],
       size=29, bold=True, align="c", spacing=4)
    pic(s, "motif.png", 3.4, 3.2, w=6.5)
    tb(s, 0.9, 5.7, 11.5, 0.9, "SETHU", size=54, bold=True, align="c", color=C_WHITE)
    tb(s, 0.9, 6.5, 11.5, 0.5, "Every message deserves a path.", size=22, color=C_CYAN, align="c", italic=True)

    return prs


if __name__ == "__main__":
    make_assets()
    build()
    out = ROOT / "Sethu_Presentation.pptx"
    prs.save(out)
    print("saved", out, "slides:", len(prs.slides))
